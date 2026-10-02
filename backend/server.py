from dotenv import load_dotenv
from pathlib import Path
import os

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

from fastapi import FastAPI, APIRouter, HTTPException, Depends, Request, Query
from fastapi.responses import StreamingResponse
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, Field, field_validator
from typing import List, Optional
from datetime import datetime, timezone, timedelta
import logging
import json
import io
import zipfile
import jwt

# MongoDB connection
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

JWT_SECRET = os.environ['JWT_SECRET']
JWT_ALGORITHM = "HS256"
AUTH_USERNAME = os.environ['AUTH_USERNAME']
AUTH_PASSWORD = os.environ['AUTH_PASSWORD']

KKTP = 75

SCHOOL = {
    "provinsi": "PEMERINTAH PROVINSI DAERAH KHUSUS IBUKOTA JAKARTA",
    "dinas": "DINAS PENDIDIKAN",
    "nama": "SMP NEGERI 37 JAKARTA",
    "alamat": "Jalan Taman Wijaya Kusuma Raya, Pondok Labu, Cilandak, Kode Pos 12450",
    "telepon": "Telepon: (021) 7695272",
    "website": "Website: www.smpn37jakarta.sch.id",
    "email": "Email: smpn37cilandak@gmail.com",
    "tahun_ajaran": "2026/2027",
    "semester": "GANJIL",
    "tanggal": "Jakarta, 9 Oktober 2026",
    "kepala_sekolah": "Triyasih, S.Pd.",
    "nip_kepala": "NIP. 196805221994122001",
}

app = FastAPI()
api_router = APIRouter(prefix="/api")

logging.basicConfig(level=logging.INFO,
                    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)


# ---------------- Auth ----------------
def create_token(username: str) -> str:
    payload = {"sub": username, "exp": datetime.now(timezone.utc) + timedelta(days=7), "type": "access"}
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGORITHM)


async def get_current_user(request: Request) -> dict:
    auth_header = request.headers.get("Authorization", "")
    token = auth_header[7:] if auth_header.startswith("Bearer ") else None
    if not token:
        raise HTTPException(status_code=401, detail="Tidak terautentikasi")
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGORITHM])
        return {"username": payload["sub"]}
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Sesi berakhir, silakan login kembali")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Token tidak valid")


class LoginInput(BaseModel):
    username: str
    password: str


@api_router.post("/auth/login")
async def login(data: LoginInput):
    if data.username != AUTH_USERNAME or data.password != AUTH_PASSWORD:
        raise HTTPException(status_code=401, detail="Username atau password salah")
    return {"token": create_token(data.username), "username": data.username}


@api_router.get("/auth/me")
async def me(user: dict = Depends(get_current_user)):
    return user


# ---------------- Models ----------------
class NilaiItem(BaseModel):
    siswa_id: str
    f1: Optional[float] = None
    f2: Optional[float] = None
    f3: Optional[float] = None
    s1: Optional[float] = None
    s2: Optional[float] = None
    s3: Optional[float] = None

    @field_validator("f1", "f2", "f3", "s1", "s2", "s3")
    @classmethod
    def clamp_range(cls, v):
        if v is None:
            return None
        return max(0.0, min(100.0, float(v)))


class WalasNipInput(BaseModel):
    nip: str = ""


class KehadiranInput(BaseModel):
    sakit: int = 0
    izin: int = 0
    alfa: int = 0
    catatan: str = ""

    @field_validator("sakit", "izin", "alfa")
    @classmethod
    def non_negative(cls, v):
        return max(0, int(v))


class KehadiranBulkItem(BaseModel):
    siswa_id: str
    sakit: int = 0
    izin: int = 0
    alfa: int = 0
    catatan: str = ""

    @field_validator("sakit", "izin", "alfa")
    @classmethod
    def non_negative(cls, v):
        return max(0, int(v))


class KehadiranBulk(BaseModel):
    kelas: str
    items: List[KehadiranBulkItem]


class NilaiBulk(BaseModel):
    kelas: str
    mapel: str
    items: List[NilaiItem]


# ---------------- Data endpoints ----------------
@api_router.get("/kelas")
async def get_kelas(user: dict = Depends(get_current_user)):
    walas = {w["kelas"]: w["nama"] async for w in db.walas.find({}, {"_id": 0})}
    kelas_list = await db.kelas.find({}, {"_id": 0}).sort("kelas", 1).to_list(1000)
    result = []
    for k in kelas_list:
        count = await db.siswa.count_documents({"kelas": k["kelas"]})
        result.append({"kelas": k["kelas"], "walas": walas.get(k["kelas"], "-"), "jumlah_siswa": count})
    return result


@api_router.get("/mapel")
async def get_mapel(user: dict = Depends(get_current_user)):
    docs = await db.mapel.find({}, {"_id": 0}).sort("urutan", 1).to_list(1000)
    return [d["nama"] for d in docs]


@api_router.get("/siswa")
async def get_siswa(kelas: str = Query(...), user: dict = Depends(get_current_user)):
    docs = await db.siswa.find({"kelas": kelas}, {"_id": 0}).sort("nama", 1).to_list(1000)
    return docs


@api_router.get("/nilai")
async def get_nilai(kelas: str = Query(...), mapel: str = Query(...), user: dict = Depends(get_current_user)):
    siswa = await db.siswa.find({"kelas": kelas}, {"_id": 0}).sort("nama", 1).to_list(1000)
    nilai_docs = await db.nilai.find({"kelas": kelas, "mapel": mapel}, {"_id": 0}).to_list(2000)
    nmap = {n["siswa_id"]: n for n in nilai_docs}
    rows = []
    for s in siswa:
        n = nmap.get(s["id"], {})
        rows.append({
            "siswa_id": s["id"], "nama": s["nama"], "nisn": s["nisn"], "jk": s["jk"],
            "f1": n.get("f1"), "f2": n.get("f2"), "f3": n.get("f3"),
            "s1": n.get("s1"), "s2": n.get("s2"), "s3": n.get("s3"),
        })
    return rows


@api_router.post("/nilai/bulk")
async def save_nilai(data: NilaiBulk, user: dict = Depends(get_current_user)):
    now = datetime.now(timezone.utc).isoformat()
    for item in data.items:
        doc = item.model_dump()
        doc.update({"kelas": data.kelas, "mapel": data.mapel, "updated_at": now})
        await db.nilai.update_one(
            {"kelas": data.kelas, "mapel": data.mapel, "siswa_id": item.siswa_id},
            {"$set": doc}, upsert=True)
    return {"status": "ok", "saved": len(data.items)}


async def _build_raport_data(kelas: str):
    mapel_docs = await db.mapel.find({}, {"_id": 0}).sort("urutan", 1).to_list(1000)
    mapel_list = [m["nama"] for m in mapel_docs]
    siswa = await db.siswa.find({"kelas": kelas}, {"_id": 0}).sort("nama", 1).to_list(1000)
    nilai_docs = await db.nilai.find({"kelas": kelas}, {"_id": 0}).to_list(20000)
    # nmap[siswa_id][mapel] = nilai
    nmap = {}
    for n in nilai_docs:
        nmap.setdefault(n["siswa_id"], {})[n["mapel"]] = n
    walas_doc = await db.walas.find_one({"kelas": kelas}, {"_id": 0})
    walas = walas_doc["nama"] if walas_doc else "-"
    walas_nip = (walas_doc or {}).get("nip", "") or ""
    keh_docs = await db.kehadiran.find({"kelas": kelas}, {"_id": 0}).to_list(2000)
    kmap = {k["siswa_id"]: k for k in keh_docs}
    return mapel_list, siswa, nmap, walas, walas_nip, kmap


@api_router.get("/walas/{kelas}")
async def get_walas(kelas: str, user: dict = Depends(get_current_user)):
    doc = await db.walas.find_one({"kelas": kelas}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Wali kelas tidak ditemukan")
    return {"kelas": kelas, "nama": doc["nama"], "nip": doc.get("nip", "") or ""}


@api_router.patch("/walas/{kelas}")
async def set_walas_nip(kelas: str, data: WalasNipInput, user: dict = Depends(get_current_user)):
    r = await db.walas.update_one({"kelas": kelas}, {"$set": {"nip": data.nip.strip()}})
    if r.matched_count == 0:
        raise HTTPException(status_code=404, detail="Wali kelas tidak ditemukan")
    return {"status": "ok"}


@api_router.get("/kehadiran/kelas/{kelas}")
async def get_kehadiran_kelas(kelas: str, user: dict = Depends(get_current_user)):
    siswa = await db.siswa.find({"kelas": kelas}, {"_id": 0}).sort("nama", 1).to_list(1000)
    docs = await db.kehadiran.find({"kelas": kelas}, {"_id": 0}).to_list(2000)
    kmap = {d["siswa_id"]: d for d in docs}
    rows = []
    for s in siswa:
        d = kmap.get(s["id"], {})
        rows.append({"siswa_id": s["id"], "nama": s["nama"], "nisn": s["nisn"],
                     "sakit": d.get("sakit", 0), "izin": d.get("izin", 0),
                     "alfa": d.get("alfa", 0), "catatan": d.get("catatan", "")})
    return rows


@api_router.post("/kehadiran/bulk")
async def save_kehadiran_bulk(data: KehadiranBulk, user: dict = Depends(get_current_user)):
    for item in data.items:
        doc = item.model_dump()
        doc.update({"kelas": data.kelas})
        await db.kehadiran.update_one({"siswa_id": item.siswa_id}, {"$set": doc}, upsert=True)
    return {"status": "ok", "saved": len(data.items)}


@api_router.get("/kehadiran/{siswa_id}")
async def get_kehadiran(siswa_id: str, user: dict = Depends(get_current_user)):
    doc = await db.kehadiran.find_one({"siswa_id": siswa_id}, {"_id": 0})
    if not doc:
        return {"siswa_id": siswa_id, "sakit": 0, "izin": 0, "alfa": 0, "catatan": ""}
    return doc


@api_router.post("/kehadiran/{siswa_id}")
async def save_kehadiran(siswa_id: str, data: KehadiranInput, user: dict = Depends(get_current_user)):
    s = await db.siswa.find_one({"id": siswa_id}, {"_id": 0})
    if not s:
        raise HTTPException(status_code=404, detail="Siswa tidak ditemukan")
    doc = data.model_dump()
    doc.update({"siswa_id": siswa_id, "kelas": s["kelas"]})
    await db.kehadiran.update_one({"siswa_id": siswa_id}, {"$set": doc}, upsert=True)
    return {"status": "ok"}


@api_router.get("/raport/student/{siswa_id}")
async def raport_student(siswa_id: str, user: dict = Depends(get_current_user)):
    s = await db.siswa.find_one({"id": siswa_id}, {"_id": 0})
    if not s:
        raise HTTPException(status_code=404, detail="Siswa tidak ditemukan")
    mapel_list, _, nmap, walas, walas_nip, kmap = await _build_raport_data(s["kelas"])
    rows = []
    sm = nmap.get(siswa_id, {})
    for m in mapel_list:
        n = sm.get(m, {})
        rows.append({"mapel": m, "kktp": KKTP,
                     "f1": n.get("f1"), "f2": n.get("f2"), "f3": n.get("f3"),
                     "s1": n.get("s1"), "s2": n.get("s2"), "s3": n.get("s3")})
    return {"siswa": s, "walas": walas, "walas_nip": walas_nip, "school": SCHOOL,
            "nilai": rows, "kehadiran": kmap.get(siswa_id, {})}


# ---------------- PDF generation ----------------
def _fmt(v):
    if v is None or v == "":
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _build_raport_story(story, styles, siswa, walas, walas_nip, mapel_rows, kehadiran):
    from reportlab.platypus import Paragraph, Spacer, Table, TableStyle, Image as RLImage
    from reportlab.lib import colors
    from reportlab.lib.units import mm

    logo_path = str(ROOT_DIR / "logo_jayaraya.png")
    header_tbl_data = [[
        RLImage(logo_path, width=22 * mm, height=22 * mm) if os.path.exists(logo_path) else "",
        Paragraph(
            f"<para align=center><font size=9>{SCHOOL['provinsi']}</font><br/>"
            f"<font size=9>{SCHOOL['dinas']}</font><br/>"
            f"<font size=14><b>{SCHOOL['nama']}</b></font><br/>"
            f"<font size=7.5>{SCHOOL['alamat']}</font><br/>"
            f"<font size=7.5>{SCHOOL['telepon']} &nbsp; {SCHOOL['website']} &nbsp; {SCHOOL['email']}</font></para>",
            styles["Normal"]),
        ""]]
    htbl = Table(header_tbl_data, colWidths=[26 * mm, 128 * mm, 26 * mm])
    htbl.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LINEBELOW", (0, 0), (-1, -1), 2, colors.black),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(htbl)
    story.append(Spacer(1, 6))
    story.append(Paragraph(
        f"<para align=center><b>LAPORAN HASIL BELAJAR TENGAH SEMESTER {SCHOOL['semester']}</b><br/>"
        f"<b>TAHUN AJARAN {SCHOOL['tahun_ajaran']}</b></para>", styles["Normal"]))
    story.append(Spacer(1, 8))

    info = Table([
        [Paragraph("Nama", styles["Normal"]), Paragraph(f": {siswa['nama']}", styles["Normal"])],
        [Paragraph("NISN", styles["Normal"]), Paragraph(f": {siswa['nisn']}", styles["Normal"])],
        [Paragraph("Kelas", styles["Normal"]), Paragraph(f": {siswa['kelas']}", styles["Normal"])],
    ], colWidths=[20 * mm, 160 * mm])
    info.setStyle(TableStyle([("BOTTOMPADDING", (0, 0), (-1, -1), 1), ("TOPPADDING", (0, 0), (-1, -1), 1)]))
    story.append(info)
    story.append(Spacer(1, 6))
    story.append(Paragraph("<b>NILAI AKADEMIK</b>", styles["Normal"]))
    story.append(Spacer(1, 3))

    cs = styles["Cell"]
    ch = styles["CellH"]
    data = [
        [Paragraph("NO", ch), Paragraph("MATA PELAJARAN", ch), Paragraph("KKTP", ch),
         Paragraph("FORMATIF", ch), "", "", Paragraph("SUMATIF", ch), "", ""],
        ["", "", "", Paragraph("F1", ch), Paragraph("F2", ch), Paragraph("F3", ch),
         Paragraph("S1", ch), Paragraph("S2", ch), Paragraph("S3", ch)],
    ]
    for i, r in enumerate(mapel_rows, 1):
        data.append([
            Paragraph(str(i), cs), Paragraph(r["mapel"], styles["CellL"]), Paragraph(str(r["kktp"]), cs),
            Paragraph(_fmt(r["f1"]), cs), Paragraph(_fmt(r["f2"]), cs), Paragraph(_fmt(r["f3"]), cs),
            Paragraph(_fmt(r["s1"]), cs), Paragraph(_fmt(r["s2"]), cs), Paragraph(_fmt(r["s3"]), cs),
        ])
    col = [10 * mm, 62 * mm, 14 * mm, 14 * mm, 14 * mm, 14 * mm, 14 * mm, 14 * mm, 14 * mm]
    tbl = Table(data, colWidths=col, repeatRows=2)
    tbl.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
        ("BACKGROUND", (0, 0), (-1, 1), colors.HexColor("#1E40AF")),
        ("TEXTCOLOR", (0, 0), (-1, 1), colors.white),
        ("SPAN", (0, 0), (0, 1)), ("SPAN", (1, 0), (1, 1)), ("SPAN", (2, 0), (2, 1)),
        ("SPAN", (3, 0), (5, 0)), ("SPAN", (6, 0), (8, 0)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 2), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
    ]))
    story.append(tbl)
    story.append(Spacer(1, 8))

    # Ketidakhadiran + catatan
    k = kehadiran or {}
    sakit, izin, alfa = int(k.get("sakit", 0) or 0), int(k.get("izin", 0) or 0), int(k.get("alfa", 0) or 0)
    catatan_text = (k.get("catatan") or "").strip()
    ket = Table([
        [Paragraph("<b>KETIDAKHADIRAN</b>", styles["Normal"]), ""],
        [Paragraph("Sakit", styles["Normal"]), Paragraph(f":  {sakit}  hari", styles["Normal"])],
        [Paragraph("Izin", styles["Normal"]), Paragraph(f":  {izin}  hari", styles["Normal"])],
        [Paragraph("Alfa", styles["Normal"]), Paragraph(f":  {alfa}  hari", styles["Normal"])],
    ], colWidths=[25 * mm, 55 * mm])
    ket.setStyle(TableStyle([("BOTTOMPADDING", (0, 0), (-1, -1), 2), ("SPAN", (0, 0), (1, 0))]))

    catatan_body = catatan_text.replace("\n", "<br/>") if catatan_text else "<br/><br/><br/>"
    catatan = Table([
        [Paragraph("<b>CATATAN WALI KELAS</b>", styles["Normal"])],
        [Paragraph(catatan_body, styles["Normal"])],
    ], colWidths=[90 * mm])
    catatan.setStyle(TableStyle([
        ("BOX", (0, 1), (-1, 1), 0.5, colors.black),
        ("TOPPADDING", (0, 1), (-1, 1), 8),
        ("BOTTOMPADDING", (0, 1), (-1, 1), 8),
        ("LEFTPADDING", (0, 1), (-1, 1), 6),
        ("RIGHTPADDING", (0, 1), (-1, 1), 6),
    ]))
    combo = Table([[ket, catatan]], colWidths=[85 * mm, 95 * mm])
    combo.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.append(combo)
    story.append(Spacer(1, 10))

    # signatures
    sig = Table([
        ["", Paragraph(f"<para align=center>{SCHOOL['tanggal']}</para>", styles["Normal"])],
        [Paragraph("<para align=center>Orangtua/Wali Murid</para>", styles["Normal"]),
         Paragraph("<para align=center>Wali Kelas</para>", styles["Normal"])],
        [Paragraph("<para align=center><br/><br/><br/>(......................)</para>", styles["Normal"]),
         Paragraph(f"<para align=center><br/><br/><br/><b>{walas}</b><br/>NIP. {walas_nip if walas_nip else '......................'}</para>", styles["Normal"])],
    ], colWidths=[90 * mm, 90 * mm])
    story.append(sig)
    story.append(Spacer(1, 6))
    story.append(Paragraph("<para align=center>Mengetahui,<br/>Kepala SMP Negeri 37 Jakarta</para>", styles["Normal"]))
    story.append(Spacer(1, 30))
    story.append(Paragraph(
        f"<para align=center><b>{SCHOOL['kepala_sekolah']}</b><br/>{SCHOOL['nip_kepala']}</para>", styles["Normal"]))


def _render_pdf(siswa_rows):
    from reportlab.platypus import SimpleDocTemplate, PageBreak
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_LEFT

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=12 * mm, bottomMargin=12 * mm,
                            leftMargin=14 * mm, rightMargin=14 * mm)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="Cell", fontName="Helvetica", fontSize=8, alignment=TA_CENTER, leading=10))
    styles.add(ParagraphStyle(name="CellL", fontName="Helvetica", fontSize=8, alignment=TA_LEFT, leading=10))
    styles.add(ParagraphStyle(name="CellH", fontName="Helvetica-Bold", fontSize=8, alignment=TA_CENTER,
                              leading=10, textColor="white"))
    styles["Normal"].fontSize = 9
    styles["Normal"].leading = 12

    story = []
    for idx, (siswa, walas, walas_nip, rows, kehadiran) in enumerate(siswa_rows):
        if idx > 0:
            story.append(PageBreak())
        _build_raport_story(story, styles, siswa, walas, walas_nip, rows, kehadiran)
    doc.build(story)
    buf.seek(0)
    return buf


@api_router.get("/raport/pdf/student/{siswa_id}")
async def raport_pdf_student(siswa_id: str, user: dict = Depends(get_current_user)):
    s = await db.siswa.find_one({"id": siswa_id}, {"_id": 0})
    if not s:
        raise HTTPException(status_code=404, detail="Siswa tidak ditemukan")
    mapel_list, _, nmap, walas, walas_nip, kmap = await _build_raport_data(s["kelas"])
    sm = nmap.get(siswa_id, {})
    rows = [{"mapel": m, "kktp": KKTP, **{k: sm.get(m, {}).get(k) for k in ["f1", "f2", "f3", "s1", "s2", "s3"]}} for m in mapel_list]
    buf = _render_pdf([(s, walas, walas_nip, rows, kmap.get(siswa_id, {}))])
    fname = f"Raport_{s['nama'].replace(' ', '_')}_{s['kelas']}.pdf"
    return StreamingResponse(buf, media_type="application/pdf",
                             headers={"Content-Disposition": f'attachment; filename="{fname}"'})


@api_router.get("/raport/pdf/kelas/{kelas}")
async def raport_pdf_kelas(kelas: str, user: dict = Depends(get_current_user)):
    mapel_list, siswa, nmap, walas, walas_nip, kmap = await _build_raport_data(kelas)
    if not siswa:
        raise HTTPException(status_code=404, detail="Kelas tidak ditemukan")
    all_rows = []
    for s in siswa:
        sm = nmap.get(s["id"], {})
        rows = [{"mapel": m, "kktp": KKTP, **{k: sm.get(m, {}).get(k) for k in ["f1", "f2", "f3", "s1", "s2", "s3"]}} for m in mapel_list]
        all_rows.append((s, walas, walas_nip, rows, kmap.get(s["id"], {})))
    buf = _render_pdf(all_rows)
    return StreamingResponse(buf, media_type="application/pdf",
                             headers={"Content-Disposition": f'attachment; filename="Raport_Kelas_{kelas}.pdf"'})


@api_router.get("/raport/zip/{kelas}")
async def raport_zip_kelas(kelas: str, user: dict = Depends(get_current_user)):
    mapel_list, siswa, nmap, walas, walas_nip, kmap = await _build_raport_data(kelas)
    if not siswa:
        raise HTTPException(status_code=404, detail="Kelas tidak ditemukan")
    zip_buf = io.BytesIO()
    with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for s in siswa:
            sm = nmap.get(s["id"], {})
            rows = [{"mapel": m, "kktp": KKTP, **{k: sm.get(m, {}).get(k) for k in ["f1", "f2", "f3", "s1", "s2", "s3"]}} for m in mapel_list]
            pdf = _render_pdf([(s, walas, walas_nip, rows, kmap.get(s["id"], {}))])
            fname = f"Raport_{s['nama'].replace(' ', '_')}_{kelas}.pdf"
            zf.writestr(fname, pdf.read())
    zip_buf.seek(0)
    return StreamingResponse(zip_buf, media_type="application/zip",
                             headers={"Content-Disposition": f'attachment; filename="Raport_{kelas}_per_siswa.zip"'})


# ---------------- Excel Leger ----------------
@api_router.get("/leger/excel/{kelas}")
async def leger_excel(kelas: str, user: dict = Depends(get_current_user)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    mapel_list, siswa, nmap, walas, walas_nip, kmap = await _build_raport_data(kelas)
    if not siswa:
        raise HTTPException(status_code=404, detail="Kelas tidak ditemukan")

    wb = Workbook()
    ws = wb.active
    ws.title = f"Leger {kelas}"

    thin = Side(style="thin", color="94A3B8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_fill = PatternFill("solid", fgColor="1E40AF")
    sub_fill = PatternFill("solid", fgColor="3B82F6")
    hdr_font = Font(bold=True, color="FFFFFF", size=10)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws["A1"] = f"LEGER NILAI KELAS {kelas} - SMP NEGERI 37 JAKARTA"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"Wali Kelas: {walas}   |   Semester {SCHOOL['semester']} T.A. {SCHOOL['tahun_ajaran']}"
    ws["A2"].font = Font(size=10)

    # Header rows start at row 4
    r1, r2 = 4, 5
    ws.cell(r1, 1, "NO"); ws.cell(r1, 2, "NISN"); ws.cell(r1, 3, "NAMA SISWA")
    ws.merge_cells(start_row=r1, start_column=1, end_row=r2, end_column=1)
    ws.merge_cells(start_row=r1, start_column=2, end_row=r2, end_column=2)
    ws.merge_cells(start_row=r1, start_column=3, end_row=r2, end_column=3)
    col = 4
    sub = ["F1", "F2", "F3", "S1", "S2", "S3"]
    for m in mapel_list:
        ws.merge_cells(start_row=r1, start_column=col, end_row=r1, end_column=col + 5)
        c = ws.cell(r1, col, m)
        c.font = hdr_font; c.fill = hdr_fill; c.alignment = center; c.border = border
        for j, s in enumerate(sub):
            sc = ws.cell(r2, col + j, s)
            sc.font = hdr_font; sc.fill = sub_fill; sc.alignment = center; sc.border = border
        col += 6
    for cc in range(1, 4):
        hc = ws.cell(r1, cc); hc.font = hdr_font; hc.fill = hdr_fill; hc.alignment = center; hc.border = border
        ws.cell(r2, cc).border = border

    row = r2 + 1
    for i, s in enumerate(siswa, 1):
        ws.cell(row, 1, i).border = border
        ws.cell(row, 2, s["nisn"]).border = border
        ws.cell(row, 3, s["nama"]).border = border
        sm = nmap.get(s["id"], {})
        col = 4
        for m in mapel_list:
            n = sm.get(m, {})
            for j, key in enumerate(["f1", "f2", "f3", "s1", "s2", "s3"]):
                v = n.get(key)
                cell = ws.cell(row, col + j, v if v is not None else "")
                cell.border = border
                cell.alignment = center
            col += 6
        row += 1

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 32
    for cidx in range(4, 4 + len(mapel_list) * 6):
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(cidx)].width = 5
    ws.freeze_panes = "D6"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="Leger_Nilai_{kelas}.xlsx"'})


# ---------------- Seeding ----------------
async def seed_data():
    seed_file = ROOT_DIR / "seed_data.json"
    if not seed_file.exists():
        logger.warning("seed_data.json not found, skipping seed")
        return
    data = json.loads(seed_file.read_text(encoding="utf-8"))

    if await db.mapel.count_documents({}) == 0:
        await db.mapel.insert_many([{"nama": m, "urutan": i} for i, m in enumerate(data["mapel"])])
    if await db.kelas.count_documents({}) == 0:
        await db.kelas.insert_many([{"kelas": k} for k in data["kelas"]])
    if await db.walas.count_documents({}) == 0:
        await db.walas.insert_many(data["walas"])
    if await db.siswa.count_documents({}) == 0:
        docs = []
        for s in data["siswa"]:
            sid = f"{s['kelas']}_{s['nisn']}"
            docs.append({"id": sid, "nama": s["nama"], "jk": s["jk"],
                         "agama": s["agama"], "nisn": s["nisn"], "kelas": s["kelas"]})
        await db.siswa.insert_many(docs)
    await db.nilai.create_index([("kelas", 1), ("mapel", 1), ("siswa_id", 1)], unique=True)
    logger.info("Seed complete: mapel=%d kelas=%d walas=%d siswa=%d",
                await db.mapel.count_documents({}), await db.kelas.count_documents({}),
                await db.walas.count_documents({}), await db.siswa.count_documents({}))


@app.on_event("startup")
async def on_startup():
    await seed_data()


app.include_router(api_router)
app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
